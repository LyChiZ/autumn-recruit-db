#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成一份虚构脱敏的 Word 示例简历（data/samples/sample_resume.docx），
供演示 DOCX 解析能力。运行：python scripts/make_sample_docx.py
"""

from __future__ import annotations

from pathlib import Path

import docx

OUT = Path(__file__).resolve().parents[1] / "data" / "samples" / "sample_resume.docx"


def main() -> None:
    doc = docx.Document()

    doc.add_paragraph("王小明 — 产品与数据分析方向（2027 届应届）")
    doc.add_paragraph("wangxiaoming@example.com ｜ 139-0000-0000")
    doc.add_paragraph("XX大学管理科学硕士在读，有产品与数据双段实习经验。")

    doc.add_heading("教育背景", level=1)
    doc.add_paragraph("XX大学 管理科学 硕士 2025.09 - 2027.06")
    doc.add_paragraph("XX大学 信息管理 本科 2021.09 - 2025.06")

    doc.add_heading("专业技能", level=1)
    doc.add_paragraph("数据分析, 产品设计, SQL, Excel, 用户研究, Axure, A/B测试")

    doc.add_heading("项目经历", level=1)
    doc.add_paragraph("- 校园外卖拼单小程序：负责需求分析与产品原型，日活峰值 2000+")
    doc.add_paragraph("- 会员体系 A/B 实验：设计实验方案并输出分析报告")

    doc.add_heading("实习经历", level=1)
    doc.add_paragraph("- 某互联网公司 产品实习生：负责 B 端后台功能迭代与数据埋点")
    doc.add_paragraph("- 某咨询公司 项目助理：参与用户访谈与案头研究")

    doc.save(str(OUT))
    print(f"已生成虚构示例：{OUT}")


if __name__ == "__main__":
    main()
