# -*- coding: utf-8 -*-
"""Word (.docx) 简历解析：基于 python-docx 抽取段落与表格文本，再交给规则抽取器。"""

from __future__ import annotations

from .base import Resume
from .extractor import extract_from_text


def parse_docx(file_path: str) -> Resume:
    """解析 .docx 简历文件。

    Args:
        file_path: .docx 文件路径

    Returns:
        Resume: 结构化简历对象
    """
    import docx  # python-docx

    document = docx.Document(file_path)

    parts: list[str] = []

    # 1. 段落文本（保留换行，便于小节切分）
    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # 2. 表格文本（简历常用表格排布）
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    full_text = "\n".join(parts)
    return extract_from_text(full_text, source_format="docx")
